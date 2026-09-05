"""Docling JSON parser for narrative families A/B/D."""

from __future__ import annotations

from collections import defaultdict
from collections.abc import Sequence
from dataclasses import replace
from pathlib import Path

from docling.chunking import HybridChunker
from docling.datamodel.base_models import ConversionStatus, InputFormat
from docling.datamodel.pipeline_options import PdfPipelineOptions, RapidOcrOptions
from docling.document_converter import (
    DocumentConverter,
    PdfFormatOption,
    PowerpointFormatOption,
    WordFormatOption,
)
from docling.pipeline.simple_pipeline import SimplePipeline
from docling.pipeline.standard_pdf_pipeline import StandardPdfPipeline
from docling_core.transforms.chunker.tokenizer.base import BaseTokenizer
from docling_core.types.doc import DocItemLabel, DoclingDocument, TableItem
from docling_core.types.doc.common.content_layer import ContentLayer
from pypdfium2 import PdfDocument, PdfiumError

from large_files_embedding.domain.document import (
    ChunkOrigin,
    ChunkType,
    DocumentFamily,
    FailureReason,
    NarrativeChunk,
    NarrativeParseError,
    ParsedNarrative,
    metadata_from_header_footer,
)


class _ApproxTokenizer(BaseTokenizer):
    max_tokens: int = 8192

    def count_tokens(self, text: str) -> int:
        return max(1, (len(text) + 1) // 2)

    def get_max_tokens(self) -> int:
        return self.max_tokens

    def get_tokenizer(self) -> object:
        return self.count_tokens


_HEADER_LABELS = {DocItemLabel.PAGE_HEADER, DocItemLabel.PAGE_FOOTER}
_HEADER_LAYERS = {
    ContentLayer.BODY,
    ContentLayer.FURNITURE,
    ContentLayer.BACKGROUND,
}


class DoclingNarrativeParser:
    def parse(
        self, path: Path, *, family: DocumentFamily, doc_id: str
    ) -> ParsedNarrative:
        if family is DocumentFamily.C:
            raise NarrativeParseError(FailureReason.TABULAR_NOT_NARRATIVE)
        needs_ocr = False
        if family is DocumentFamily.B:
            needs_ocr = self._pdf_needs_ocr(path)
        try:
            converter = self._converter(family, needs_ocr=needs_ocr)
            result = converter.convert(path)
        except NarrativeParseError:
            raise
        except Exception as exc:
            raise _map_error(exc) from exc
        if result.status is ConversionStatus.FAILURE:
            raise NarrativeParseError(FailureReason.PARSE_FAILED)
        document = result.document
        json_bytes = document.model_dump_json().encode("utf-8")
        headers = _header_footer_texts(document)
        if family is DocumentFamily.D:
            chunks = self._slide_chunks(document, path, doc_id)
            pipeline = "slide"
        else:
            chunks = self._hybrid_chunks(
                document, path, family, doc_id, headers=headers
            )
            if family is DocumentFamily.A:
                pipeline = "simple"
            elif needs_ocr:
                pipeline = "rapid_ocr"
            else:
                pipeline = "standard_pdf"
        header_footer = "\n".join(headers) or None
        part_no, vehicle = metadata_from_header_footer(header_footer or "")
        chunks = [
            _with_header_meta(chunk, headers=headers, part_no=part_no, vehicle=vehicle)
            for chunk in chunks
        ]
        chunks = [
            chunk
            for chunk in chunks
            if chunk.text.strip() or chunk.embedding_input.strip()
        ]
        if not chunks:
            raise NarrativeParseError(FailureReason.EMPTY_DOCUMENT)
        ocr_lang = ("korean",) if needs_ocr else None
        return ParsedNarrative(
            doc_id=doc_id,
            path=path,
            family=family,
            json_bytes=json_bytes,
            chunks=tuple(chunks),
            pipeline=pipeline,
            ocr_lang=ocr_lang,
            header_footer=header_footer,
        )

    def _converter(
        self, family: DocumentFamily, *, needs_ocr: bool
    ) -> DocumentConverter:
        if family is DocumentFamily.A:
            return DocumentConverter(
                allowed_formats=[InputFormat.DOCX],
                format_options={
                    InputFormat.DOCX: WordFormatOption(pipeline_cls=SimplePipeline),
                },
            )
        if family is DocumentFamily.D:
            return DocumentConverter(
                allowed_formats=[InputFormat.PPTX],
                format_options={
                    InputFormat.PPTX: PowerpointFormatOption(
                        pipeline_cls=SimplePipeline
                    ),
                },
            )
        options = PdfPipelineOptions()
        options.do_ocr = needs_ocr
        options.do_table_structure = True
        options.do_chart_extraction = False
        options.ocr_options = RapidOcrOptions(lang=["korean"])
        return DocumentConverter(
            allowed_formats=[InputFormat.PDF],
            format_options={
                InputFormat.PDF: PdfFormatOption(
                    pipeline_cls=StandardPdfPipeline,
                    pipeline_options=options,
                ),
            },
        )

    def _pdf_needs_ocr(self, path: Path) -> bool:
        try:
            pdf = PdfDocument(str(path))
        except PdfiumError as exc:
            raise _map_error(exc) from exc
        except Exception as exc:
            raise _map_error(exc) from exc
        try:
            page_count = len(pdf)
            if page_count <= 0:
                raise NarrativeParseError(FailureReason.EMPTY_DOCUMENT)
            sample_at = sorted({0, page_count // 2, page_count - 1})
            chars = 0
            for index in sample_at:
                page = pdf[index]
                textpage = page.get_textpage()
                try:
                    chars += len(textpage.get_text_range().strip())
                finally:
                    textpage.close()
            return chars < 40 * len(sample_at)
        finally:
            pdf.close()

    def _hybrid_chunks(
        self,
        document: DoclingDocument,
        path: Path,
        family: DocumentFamily,
        doc_id: str,
        *,
        headers: Sequence[str] = (),
    ) -> list[NarrativeChunk]:
        chunker = HybridChunker(
            tokenizer=_ApproxTokenizer(),
            merge_peers=True,
            repeat_table_header=True,
        )
        chunks: list[NarrativeChunk] = []
        extra_headers = list(headers)
        for index, raw in enumerate(chunker.chunk(dl_doc=document)):
            if _header_footer_only(raw):
                extra = (raw.text or "").strip()
                if extra and extra not in extra_headers:
                    extra_headers.append(extra)
                continue
            text = _strip_headers(raw.text or "", extra_headers)
            if not text.strip():
                continue
            embedding_input = _strip_headers(chunker.contextualize(raw), extra_headers)
            headings = list(getattr(raw.meta, "headings", None) or [])
            section_path = " > ".join(headings) if headings else None
            chunk_type = ChunkType.TABLE if _has_table(raw) else ChunkType.TEXT
            chunks.append(
                NarrativeChunk(
                    chunk_id=f"{doc_id}-{index}",
                    doc_id=doc_id,
                    path=str(path),
                    family=family,
                    chunk_type=chunk_type,
                    text=text,
                    embedding_input=embedding_input,
                    origin=ChunkOrigin.JSON,
                    section_path=section_path,
                    parent_id=(f"{doc_id}:{section_path}" if section_path else None),
                    page=_page_no(raw),
                )
            )
        return chunks

    def _slide_chunks(
        self, document: DoclingDocument, path: Path, doc_id: str
    ) -> list[NarrativeChunk]:
        notes = _pptx_notes(path)
        by_page: dict[int, list[object]] = defaultdict(list)
        for item, _level in document.iterate_items():
            label = getattr(item, "label", None)
            if label in _HEADER_LABELS:
                continue
            page_no = _item_page(item)
            if page_no is None:
                continue
            by_page[page_no].append(item)
        chunks: list[NarrativeChunk] = []
        index = 0
        for slide_index in sorted(by_page):
            titles: list[str] = []
            body: list[str] = []
            tables: list[object] = []
            for node in by_page[slide_index]:
                label = getattr(node, "label", None)
                if label is DocItemLabel.TABLE or isinstance(node, TableItem):
                    tables.append(node)
                    continue
                text = (getattr(node, "text", None) or "").strip()
                if not text:
                    continue
                if label is DocItemLabel.TITLE:
                    titles.append(text)
                else:
                    body.append(text)
            note = notes.get(slide_index, "")
            title = " ".join(titles)
            body_text = "\n".join(body)
            parts = [f"[슬라이드 {slide_index}]"]
            if title:
                parts.append(title)
            if body_text:
                parts.append(body_text)
            if note:
                parts.append(f"노트: {note}")
            embedding_input = "\n".join(parts)
            visible = "\n".join(part for part in (title, body_text) if part)
            if visible.strip() or note:
                chunks.append(
                    NarrativeChunk(
                        chunk_id=f"{doc_id}-{index}",
                        doc_id=doc_id,
                        path=str(path),
                        family=DocumentFamily.D,
                        chunk_type=ChunkType.TEXT,
                        text=visible or embedding_input,
                        embedding_input=embedding_input,
                        origin=ChunkOrigin.JSON,
                        slide_index=slide_index,
                        page=slide_index,
                        speaker_notes=note or None,
                    )
                )
                index += 1
            for table in tables:
                table_text = _table_text(table)
                if not table_text.strip():
                    continue
                table_input = f"[슬라이드 {slide_index}]\n{table_text}"
                chunks.append(
                    NarrativeChunk(
                        chunk_id=f"{doc_id}-{index}",
                        doc_id=doc_id,
                        path=str(path),
                        family=DocumentFamily.D,
                        chunk_type=ChunkType.TABLE,
                        text=table_text,
                        embedding_input=table_input,
                        origin=ChunkOrigin.JSON,
                        slide_index=slide_index,
                        page=slide_index,
                        speaker_notes=None,
                    )
                )
                index += 1
        return chunks


def _pptx_notes(path: Path) -> dict[int, str]:
    try:
        from pptx import Presentation
    except ImportError:
        return {}
    try:
        presentation = Presentation(str(path))
    except Exception:
        return {}
    notes: dict[int, str] = {}
    for slide_index, slide in enumerate(presentation.slides, start=1):
        if not slide.has_notes_slide:
            continue
        text = slide.notes_slide.notes_text_frame.text
        if text and text.strip():
            notes[slide_index] = text.strip()
    return notes


def _header_footer_texts(document: DoclingDocument) -> list[str]:
    found: list[str] = []
    for item, _level in document.iterate_items(included_content_layers=_HEADER_LAYERS):
        if getattr(item, "label", None) not in _HEADER_LABELS:
            continue
        text = (getattr(item, "text", None) or "").strip()
        if text and text not in found:
            found.append(text)
    return found


def _strip_headers(text: str, headers: Sequence[str]) -> str:
    result = text
    for header in sorted((h for h in headers if h), key=len, reverse=True):
        result = result.replace(header, "")
    return result.strip()


def _with_header_meta(
    chunk: NarrativeChunk,
    *,
    headers: Sequence[str],
    part_no: str | None,
    vehicle: str | None,
) -> NarrativeChunk:
    return replace(
        chunk,
        text=_strip_headers(chunk.text, headers),
        embedding_input=_strip_headers(chunk.embedding_input, headers),
        part_no=chunk.part_no or part_no,
        vehicle=chunk.vehicle or vehicle,
    )


def _header_footer_only(chunk: object) -> bool:
    items = list(getattr(getattr(chunk, "meta", None), "doc_items", None) or [])
    if not items:
        return False
    labels = [getattr(item, "label", None) for item in items]
    return bool(labels) and all(label in _HEADER_LABELS for label in labels)


def _has_table(chunk: object) -> bool:
    items = list(getattr(getattr(chunk, "meta", None), "doc_items", None) or [])
    return any(
        getattr(item, "label", None) is DocItemLabel.TABLE
        or isinstance(item, TableItem)
        for item in items
    )


def _page_no(chunk: object) -> int | None:
    items = list(getattr(getattr(chunk, "meta", None), "doc_items", None) or [])
    for item in items:
        page = _item_page(item)
        if page is not None:
            return page
    return None


def _item_page(item: object) -> int | None:
    prov = getattr(item, "prov", None) or []
    if not prov:
        return None
    return getattr(prov[0], "page_no", None)


def _table_text(item: object) -> str:
    export = getattr(item, "export_to_markdown", None)
    if callable(export):
        try:
            return str(export())
        except Exception:
            pass
    return str(getattr(item, "text", "") or "")


def _map_error(exc: BaseException) -> NarrativeParseError:
    message = str(exc).lower()
    if "password" in message or "encrypt" in message:
        return NarrativeParseError(FailureReason.ENCRYPTED_PDF)
    if "xref" in message:
        return NarrativeParseError(FailureReason.BROKEN_XREF)
    if "empty" in message:
        return NarrativeParseError(FailureReason.EMPTY_DOCUMENT)
    return NarrativeParseError(FailureReason.PARSE_FAILED)
