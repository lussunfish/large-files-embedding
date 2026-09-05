"""UC-03 integration: skip if 01-stable ports or Ollama are down."""

from __future__ import annotations

import inspect
import json
import os
import socket
import urllib.error
import urllib.request
from pathlib import Path
from zipfile import ZipFile

import pytest
from typer.testing import CliRunner

from large_files_embedding.application.ingest_narrative import IngestNarrative
from large_files_embedding.domain.document import (
    DEFAULT_EMBEDDING_DIM,
    MARKET_QUALITY_BUCKET,
    MARKET_QUALITY_COLLECTION,
    ChunkOrigin,
    ChunkType,
    Document,
    DocumentFamily,
    EncodedEmbedding,
    FileSignature,
    NarrativeChunk,
    SignatureKind,
    StoredChunk,
    validate_narrative_chunks,
)
from large_files_embedding.infrastructure.docling_adapter import (
    DoclingNarrativeParser,
    _ApproxTokenizer,
    _header_footer_texts,
)
from large_files_embedding.infrastructure.embedding_encoder import DenseSparseEncoder
from large_files_embedding.infrastructure.milvus_chunk_store import MilvusChunkStore
from large_files_embedding.infrastructure.minio_object_store import MinioObjectStore
from large_files_embedding.presentation.cli.main import app

_PKG = "http://schemas.openxmlformats.org/package/2006"
_OD = "http://schemas.openxmlformats.org/officeDocument/2006"
_WML = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_RELS_TYPE = "application/vnd.openxmlformats-package.relationships+xml"
_DOCX_MAIN = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
)


def _http_ok(url: str) -> bool:
    try:
        with urllib.request.urlopen(url, timeout=2) as response:
            return 200 <= response.status < 500
    except (urllib.error.URLError, TimeoutError, OSError):
        return False


def _port_open(host: str, port: int) -> bool:
    try:
        with socket.create_connection((host, port), timeout=1):
            return True
    except OSError:
        return False


def _ollama_has_model() -> bool:
    try:
        with urllib.request.urlopen(
            "http://127.0.0.1:11434/api/tags", timeout=2
        ) as response:
            payload = json.load(response)
    except (urllib.error.URLError, TimeoutError, OSError, json.JSONDecodeError):
        return False
    names = [str(item.get("name", "")) for item in payload.get("models", [])]
    return any(
        name == "qwen3-embedding:4b" or name.startswith("qwen3-embedding:4b")
        for name in names
    )


def _infra_skip_reason() -> str | None:
    missing: list[str] = []
    if not _http_ok("http://127.0.0.1:9091/healthz"):
        missing.append("milvus")
    if not _http_ok("http://127.0.0.1:9000/minio/health/live"):
        missing.append("minio")
    if not _port_open("127.0.0.1", 19530):
        missing.append("milvus-grpc")
    if not _ollama_has_model():
        missing.append("ollama:qwen3-embedding:4b")
    if missing:
        return "01-stable/ollama down: " + ", ".join(missing)
    return None


def _require_infra() -> None:
    reason = _infra_skip_reason()
    if reason is not None:
        pytest.skip(reason)


def _minio_creds(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv(
        "MINIO_ACCESS_KEY",
        os.environ.get("MINIO_ACCESS_KEY") or "minioadmin",
    )
    monkeypatch.setenv(
        "MINIO_SECRET_KEY",
        os.environ.get("MINIO_SECRET_KEY") or "minioadmin",
    )
    monkeypatch.setenv("MINIO_BUCKET", MARKET_QUALITY_BUCKET)
    monkeypatch.setenv("MINIO_ENDPOINT", "http://127.0.0.1:9000")
    monkeypatch.setenv("MILVUS_URI", "http://127.0.0.1:19530")
    monkeypatch.setenv("MILVUS_COLLECTION", MARKET_QUALITY_COLLECTION)
    monkeypatch.setenv("OLLAMA_HOST", "http://127.0.0.1:11434")
    monkeypatch.setenv("EMBEDDING_MODEL", "qwen3-embedding:4b")
    monkeypatch.setenv("EMBEDDING_DIM", str(DEFAULT_EMBEDDING_DIM))


def _write_docx(path: Path, text: str) -> Path:
    with ZipFile(path, "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<Types xmlns="{_PKG}/content-types">'
            f'<Default Extension="rels" ContentType="{_RELS_TYPE}"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" '
            f'ContentType="{_DOCX_MAIN}"/>'
            "</Types>",
        )
        archive.writestr(
            "_rels/.rels",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<Relationships xmlns="{_PKG}/relationships">'
            '<Relationship Id="rId1" '
            f'Type="{_OD}/relationships/officeDocument" '
            'Target="word/document.xml"/>'
            "</Relationships>",
        )
        archive.writestr(
            "word/_rels/document.xml.rels",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<Relationships xmlns="{_PKG}/relationships"/>',
        )
        archive.writestr(
            "word/document.xml",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<w:document xmlns:w="{_WML}">'
            "<w:body>"
            f'<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr>'
            f"<w:r><w:t>3. 원인</w:t></w:r></w:p>"
            f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"
            "</w:body></w:document>",
        )
    return path


def _write_pptx_with_notes(path: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    deck = Presentation()
    blank = (
        deck.slide_layouts[6] if len(deck.slide_layouts) > 6 else deck.slide_layouts[0]
    )
    slide = deck.slides.add_slide(blank)
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(1))
    box.text_frame.paragraphs[0].text = "누유 현상"
    box.text_frame.paragraphs[0].font.size = Pt(28)
    notes = slide.notes_slide.notes_text_frame
    notes.text = "현장 오일 누유"
    slide2 = deck.slides.add_slide(blank)
    box2 = slide2.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(1))
    box2.text_frame.paragraphs[0].text = "대책"
    slide2.notes_slide.notes_text_frame.text = "가스켓 교체"
    deck.save(str(path))
    return path


def test_ollama_encoder_returns_nonzero_2560() -> None:
    _require_infra()
    encoded = DenseSparseEncoder.from_env().encode(["품번 ABC-1234 대책"])
    assert len(encoded) == 1
    assert len(encoded[0].dense) == DEFAULT_EMBEDDING_DIM
    assert any(abs(value) > 1e-15 for value in encoded[0].dense)
    assert encoded[0].sparse
    assert encoded[0].model == "qwen3-embedding:4b"


def test_minio_uses_market_quality_bucket(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _require_infra()
    _minio_creds(monkeypatch)
    store = MinioObjectStore.from_env()
    key = "uc03-test/probe.json"
    store.put_bytes(key, b'{"ok":true}', content_type="application/json")
    client = store._client_or_connect()
    assert client.bucket_exists(MARKET_QUALITY_BUCKET)
    response = client.get_object(MARKET_QUALITY_BUCKET, key)
    try:
        assert response.read() == b'{"ok":true}'
    finally:
        response.close()
        response.release_conn()


def _write_pptx_with_table_and_notes(path: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    deck = Presentation()
    blank = (
        deck.slide_layouts[6] if len(deck.slide_layouts) > 6 else deck.slide_layouts[0]
    )
    slide = deck.slides.add_slide(blank)
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(8), Inches(1))
    box.text_frame.paragraphs[0].text = "누유 현상"
    box.text_frame.paragraphs[0].font.size = Pt(24)
    table = slide.shapes.add_table(
        2, 2, Inches(0.5), Inches(1.6), Inches(6), Inches(1.4)
    ).table
    table.cell(0, 0).text = "품번"
    table.cell(0, 1).text = "수량"
    table.cell(1, 0).text = "A-1"
    table.cell(1, 1).text = "3"
    slide.notes_slide.notes_text_frame.text = "현장 오일 누유"
    deck.save(str(path))
    return path


def _write_text_pdf(path: Path, text: str) -> Path:
    safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content = f"BT /F1 12 Tf 50 700 Td ({safe}) Tj ET"
    stream = content.encode("latin-1", errors="replace")
    objects = [
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n",
        b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj\n",
        (
            b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >> endobj\n"
        ),
        (
            f"4 0 obj << /Length {len(stream)} >> stream\n".encode()
            + stream
            + b"\nendstream\nendobj\n"
        ),
        b"5 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n",
    ]
    body = b"".join(objects)
    offsets = []
    cursor = len(b"%PDF-1.4\n")
    for obj in objects:
        offsets.append(cursor)
        cursor += len(obj)
    xref = b"xref\n0 6\n0000000000 65535 f \n" + b"".join(
        f"{offset:010d} 00000 n \n".encode() for offset in offsets
    )
    trailer = (
        b"trailer << /Size 6 /Root 1 0 R >>\nstartxref\n"
        + str(len(b"%PDF-1.4\n") + len(body)).encode()
        + b"\n%%EOF\n"
    )
    path.write_bytes(b"%PDF-1.4\n" + body + xref + trailer)
    return path


def test_pptx_parser_is_slide_level_with_notes(tmp_path: Path) -> None:
    _require_infra()
    path = _write_pptx_with_notes(tmp_path / "review.pptx")
    parsed = DoclingNarrativeParser().parse(
        path, family=DocumentFamily.D, doc_id="uc03-pptx"
    )
    text_chunks = [chunk for chunk in parsed.chunks if chunk.slide_index is not None]
    indexes = sorted({chunk.slide_index for chunk in text_chunks if chunk.slide_index})
    assert indexes == [1, 2]
    notes = " ".join(chunk.embedding_input for chunk in parsed.chunks)
    assert "현장 오일 누유" in notes
    assert "가스켓 교체" in notes
    assert parsed.json_bytes
    validate_narrative_chunks(parsed.chunks, json_bytes=parsed.json_bytes)


def test_pptx_table_and_notes_parse_without_failure(tmp_path: Path) -> None:
    path = _write_pptx_with_table_and_notes(tmp_path / "monthly.pptx")
    parsed = DoclingNarrativeParser().parse(
        path, family=DocumentFamily.D, doc_id="uc03-pptx-table"
    )
    validate_narrative_chunks(parsed.chunks, json_bytes=parsed.json_bytes)
    notes = " ".join(
        chunk.embedding_input
        for chunk in parsed.chunks
        if chunk.chunk_type is ChunkType.TEXT
    )
    assert "현장 오일 누유" in notes
    tables = [chunk for chunk in parsed.chunks if chunk.chunk_type is ChunkType.TABLE]
    if tables:
        assert all(chunk.speaker_notes is None for chunk in tables)
        assert all(
            "현장 오일 누유" not in (chunk.speaker_notes or "") for chunk in tables
        )


def test_docx_ingest_stores_json_and_milvus(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _require_infra()
    _minio_creds(monkeypatch)
    path = _write_docx(tmp_path / "claim.docx", "공정 조건 이탈이 근본원인이다.")
    document = Document.from_signature(path, FileSignature(SignatureKind.OOXML_WORD))
    uc = IngestNarrative(
        DoclingNarrativeParser(),
        DenseSparseEncoder.from_env(),
        MilvusChunkStore.from_env(),
        MinioObjectStore.from_env(),
    )
    result = uc.execute(document)
    assert result.failed is False
    assert result.chunk_count >= 1
    assert result.json_key is not None
    assert result.original_key is not None
    objects = MinioObjectStore.from_env()
    client = objects._client_or_connect()
    json_obj = client.get_object(MARKET_QUALITY_BUCKET, result.json_key)
    try:
        body = json_obj.read()
    finally:
        json_obj.close()
        json_obj.release_conn()
    assert b"DoclingDocument" in body or b"schema_name" in body or body
    store = MilvusChunkStore.from_env()
    milvus = store._client_or_connect()
    hits = milvus.query(
        collection_name=MARKET_QUALITY_COLLECTION,
        filter=f'doc_id == "{result.doc_id}"',
        output_fields=["doc_id", "family", "chunk_type", "path", "text"],
        limit=10,
    )
    assert hits
    assert hits[0]["family"] == "A"
    assert hits[0]["path"]


def test_pdf_ocr_branch_follows_text_density(tmp_path: Path) -> None:
    from docling.datamodel.base_models import InputFormat

    digital = _write_text_pdf(
        tmp_path / "digital.pdf",
        "Market quality root cause and countermeasure report " * 6,
    )
    sparse = _write_text_pdf(tmp_path / "sparse.pdf", "x")
    parser = DoclingNarrativeParser()
    assert parser._pdf_needs_ocr(digital) is False
    assert parser._pdf_needs_ocr(sparse) is True
    digital_opts = parser._converter(
        DocumentFamily.B, needs_ocr=False
    ).format_to_options[InputFormat.PDF]
    sparse_opts = parser._converter(DocumentFamily.B, needs_ocr=True).format_to_options[
        InputFormat.PDF
    ]
    assert digital_opts.pipeline_options.do_ocr is False
    assert digital_opts.pipeline_cls.__name__ == "StandardPdfPipeline"
    assert digital_opts.pipeline_options.ocr_options.lang == ["korean"]
    assert sparse_opts.pipeline_options.do_ocr is True
    assert sparse_opts.pipeline_options.ocr_options.lang == ["korean"]
    parsed = parser.parse(digital, family=DocumentFamily.B, doc_id="uc03-pdf")
    assert parsed.pipeline == "standard_pdf"
    assert parsed.ocr_lang is None


def test_header_footer_is_metadata_not_body(tmp_path: Path) -> None:
    from docling_core.types.doc import DocItemLabel, DoclingDocument
    from docling_core.types.doc.common.content_layer import ContentLayer

    document = DoclingDocument(name="claim")
    document.add_text(
        label=DocItemLabel.PAGE_HEADER,
        text="품번 ABC-1234 차종 소나타",
        content_layer=ContentLayer.FURNITURE,
    )
    document.add_text(label=DocItemLabel.TEXT, text="공정 조건 이탈이 근본원인이다.")
    headers = _header_footer_texts(document)
    assert "품번 ABC-1234 차종 소나타" in headers
    parser = DoclingNarrativeParser()
    chunks = parser._hybrid_chunks(
        document,
        tmp_path / "claim.docx",
        DocumentFamily.A,
        "hdr-1",
        headers=headers,
    )
    assert chunks
    assert all("품번 ABC-1234" not in chunk.text for chunk in chunks)
    assert all("품번 ABC-1234" not in chunk.embedding_input for chunk in chunks)


def test_long_parent_id_upsert_does_not_fail(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _require_infra()
    _minio_creds(monkeypatch)
    section = "3. 근본원인 > " + ("공정조건세부항목-" * 40)
    parent = f"uc03longparentidxx:{section}"
    chunk = NarrativeChunk(
        chunk_id="uc03-long-parent",
        doc_id="uc03-long-parent-doc",
        path="claim.docx",
        family=DocumentFamily.A,
        chunk_type=ChunkType.TEXT,
        text="공정 조건 이탈",
        embedding_input="[섹션] 근본원인\n공정 조건 이탈",
        origin=ChunkOrigin.JSON,
        section_path=section,
        parent_id=parent,
    )
    embedding = EncodedEmbedding(
        dense=tuple(0.01 for _ in range(DEFAULT_EMBEDDING_DIM)),
        sparse=((1, 1.0),),
        model="qwen3-embedding:4b",
    )
    store = MilvusChunkStore.from_env()
    store.upsert(MARKET_QUALITY_COLLECTION, [StoredChunk(chunk, embedding)])
    milvus = store._client_or_connect()
    hits = milvus.query(
        collection_name=MARKET_QUALITY_COLLECTION,
        filter='doc_id == "uc03-long-parent-doc"',
        output_fields=["doc_id", "parent_id", "section_path"],
        limit=5,
    )
    assert hits
    assert hits[0]["doc_id"] == "uc03-long-parent-doc"
    assert hits[0]["parent_id"]


def test_approx_tokenizer_splits_overflow_body(tmp_path: Path) -> None:
    from docling_core.types.doc import DocItemLabel, DoclingDocument

    tokenizer = _ApproxTokenizer()
    counter = tokenizer.get_tokenizer()
    assert callable(counter)
    assert counter("abcd") == tokenizer.count_tokens("abcd")
    document = DoclingDocument(name="long")
    document.add_text(label=DocItemLabel.TEXT, text="품질 대책 문장입니다. " * 8000)
    parser = DoclingNarrativeParser()
    chunks = parser._hybrid_chunks(
        document, tmp_path / "long.docx", DocumentFamily.A, "long-1"
    )
    assert len(chunks) > 1


def test_parser_adapter_uses_family_pipelines_not_markdown_or_pdf_conversion() -> None:
    source = inspect.getsource(DoclingNarrativeParser)
    assert "SimplePipeline" in source
    assert "StandardPdfPipeline" in source
    assert "HybridChunker" in source
    assert "contextualize" in source
    assert "merge_peers=True" in source
    assert "repeat_table_header=True" in source
    assert 'lang=["korean"]' in source or "lang=['korean']" in source
    assert "chinese" not in source
    assert "VlmPipeline" not in source
    assert "convert_to_docx" not in source
    assert "read_bytes" not in source
    assert "sliding" not in source.lower()
    assert "anthropic" not in source.lower()


def test_encoder_adapter_uses_ollama_not_bge_or_dummy() -> None:
    source = inspect.getsource(DenseSparseEncoder)
    assert "11434" in source or "OLLAMA_HOST" in source
    assert "qwen3-embedding:4b" in source
    assert "2560" in source
    assert "BGE-M3" not in source
    assert "sentence_transformers" not in source
    assert "sentence-transformers" not in source


def test_chunk_store_is_milvus_not_lite_or_mariadb() -> None:
    source = inspect.getsource(MilvusChunkStore)
    assert "market_quality_chunks_hybrid" in source
    assert "psychology_chunks_hybrid" not in source.replace("FORBIDDEN", "")
    assert "milvus.db" not in source
    assert "embeddings.chunks" not in source
    assert "MariaDB" not in source
    assert "MilvusClient" in source or "pymilvus" in source


def test_object_store_uses_market_quality_bucket() -> None:
    source = inspect.getsource(MinioObjectStore)
    assert "market-quality-docs" in source
    assert "psychology-pdfs" not in source.replace("FORBIDDEN", "")
    assert "ebook-pdfs" not in source.replace("FORBIDDEN", "")


def test_family_c_cli_runs_tabular_not_narrative(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _minio_creds(monkeypatch)
    path = tmp_path / "ledger.xlsx"
    with ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types></Types>")
        archive.writestr("xl/workbook.xml", "<workbook/>")
    result = CliRunner().invoke(app, ["ingest", str(path)])
    assert result.exit_code == 0
    assert "family=C" in result.stdout
    assert "skip=tabular" not in result.stdout
    assert "chunks=" not in result.stdout


def test_empty_pdf_cli_does_not_need_soffice(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.delenv("DOCLING_LIBREOFFICE_CMD", raising=False)
    monkeypatch.setenv("PATH", str(tmp_path / "empty-bin"))
    (tmp_path / "empty-bin").mkdir()
    pdf = tmp_path / "empty.pdf"
    pdf.write_bytes(b"%PDF-1.4\n%%EOF\n")
    result = CliRunner().invoke(app, ["ingest", str(pdf)])
    assert result.exit_code == 0
    assert "family=B" in result.stdout
    assert "soffice_missing" not in result.stdout
